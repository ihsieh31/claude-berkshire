"""匯出（M6e）：CSV、JSON、XLSX。資料集與儀表板同源（metrics），不另算。"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

from . import metrics as M

DATASETS = ("holdings", "orders", "fills", "cash", "corporate_actions", "analyses", "nav", "usage", "approvals",
            "anomalies", "registry")
SCHEMA_VERSION = 2   # 匯出格式版本：2＝anomalies 含已結案列、加 registry 與 meta（M8-D5-02、M8-D5-08）
FORMATS = ("csv", "json", "xlsx")
# 資料集說明（寫進 all.json／xlsx 的 meta）：名稱與內容不直覺的才列（A3-M8C2-09）
NOTES = {"registry": "登錄表非持股列（觀察名單與排除名單）；持股見 holdings",
         "analyses": "未作廢的決策；provenance＝來源（NULL＝本機原生，他機匯入由 hark/export.sh 填）"}


def _flat(v: Any) -> Any:
    return json.dumps(v, ensure_ascii=False, default=str) if isinstance(v, (dict, list)) else v


def dataset(conn, name: str) -> list[dict]:
    if name not in DATASETS:
        raise ValueError(f"資料集須為 {'、'.join(DATASETS)}")
    if name == "holdings":
        return M.holdings(conn)
    if name in ("orders", "fills", "cash", "corporate_actions"):
        return M.trades(conn)[name]
    if name == "analyses":
        return M.analyses(conn, limit=100000)
    if name == "nav":
        rows = {r["date"]: dict(r) for r in conn.execute(
            "SELECT date, nav, cash, positions_value, n_positions, benchmark_close FROM nav_daily")}
        return [{**rows[x["date"]], "account_index": x["account_index"], "spy_index": x["spy_index"]}
                for x in M.series(conn)]
    if name == "usage":
        return M.usage(conn)
    if name == "approvals":
        return [dict(r) for r in conn.execute("SELECT * FROM approvals ORDER BY approval_id")]
    if name == "registry":   # 非持股（觀察／排除名單）；持股在 holdings（A3-M8C2-09）
        return M.watchlist(conn)
    # 全部異常（含已結案與 manual_adjust 稽核列；status／resolved_at／resolution 照表，M8-D5-08）
    return [{**dict(r), "hint": M.anomaly_hint(dict(r))} for r in conn.execute("SELECT * FROM anomalies ORDER BY anomaly_id")]


def meta(conn) -> dict:
    """all.json／xlsx 的 meta：匯出時間、淨值資料日、格式版本、資料集說明。"""
    from .. import db
    nav = conn.execute("SELECT MAX(date) FROM nav_daily").fetchone()[0]
    return {"exported_at": db.now_iso(), "nav_as_of": nav, "schema_version": SCHEMA_VERSION, "notes": NOTES}


def _columns(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def to_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    cols = _columns(rows)
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({k: _flat(r.get(k)) for k in cols})
    return "﻿" + buf.getvalue()  # BOM：Excel 開中文不亂碼


def _xlsx_cell(v: Any) -> Any:
    """字串去掉 xlsx 不容許的控制字元（否則整份匯出 500）；數字等照原樣（A3-M8E-05）。"""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    v = _flat(v)
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def to_xlsx(data: dict[str, list[dict]]) -> bytes:
    """每個資料集一張工作表。以 = 開頭的字串存成文字、不當公式（模型理由等外來文字，A3-M8E-05）。"""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in data.items():
        ws = wb.create_sheet(name[:31])
        cols = _columns(rows)
        ws.append([_xlsx_cell(k) for k in cols])
        for r in rows:
            ws.append([_xlsx_cell(r.get(k)) for k in cols])
            for c in ws[ws.max_row]:
                if c.data_type == "f":
                    c.data_type = "s"
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_json(data: dict[str, list[dict]]) -> str:
    return json.dumps(data, ensure_ascii=False, indent=1, default=str)


def export(conn, names: list[str], fmt: str, out_dir: Path, stamp: str) -> list[Path]:
    """寫檔：csv 每個資料集一檔；json、xlsx 合成一檔。回傳寫出的路徑。"""
    if fmt not in FORMATS:
        raise ValueError(f"格式須為 {'、'.join(FORMATS)}")
    data = {n: dataset(conn, n) for n in names}
    if fmt != "csv":
        data = {"meta": [meta(conn)], **data}
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    if fmt == "csv":
        paths = []
        for n, rows in data.items():
            p = out_dir / f"{n}_{stamp}.csv"
            p.write_text(to_csv(rows), encoding="utf-8")
            paths.append(p)
        return paths
    p = out_dir / f"berkshire_{stamp}.{fmt}"
    if fmt == "json":
        p.write_text(to_json(data), encoding="utf-8")
    else:
        p.write_bytes(to_xlsx(data))
    return [p]
