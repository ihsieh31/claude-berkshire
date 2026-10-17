"""M6c–e：網頁安全標頭、匯出、前瞻檢討、launchd 設定。不連網。"""
from __future__ import annotations

import io
import json
import plistlib

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from test_approvals import _env
from test_dashboard import H, PORT, _book

from trader import cli, db
from trader.dashboard import launchd as L
from trader.dashboard.server import create_app
from trader.report import export as X
from trader.report import validation as V


class FakeAlpaca:
    def __init__(self, closes):  # {ticker: (首, 末)}
        self.closes = closes

    def bars(self, symbols, start, end, adjustment="raw"):
        assert adjustment == "all"
        rows = []
        for t in symbols:
            if t in self.closes:
                a, b = self.closes[t]
                rows += [{"ticker": t, "date": start, "close": a}, {"ticker": t, "date": end, "close": b}]
        return pd.DataFrame(rows, columns=["ticker", "date", "close"])


def test_index_has_csp_and_static(tmp_path):
    env = _env(tmp_path)
    c = TestClient(create_app(env.paths, PORT, spawn=lambda: None), base_url=f"http://127.0.0.1:{PORT}")
    r = c.get("/", headers=H)
    assert "script-src 'self'" in r.headers["content-security-policy"]
    assert '<script src="/static/app.js">' in r.text and "{{DASHBOARD_TOKEN}}" not in r.text
    assert c.get("/static/app.js", headers=H).status_code == 200
    assert c.get("/static/app.css", headers=H).status_code == 200


def test_export_formats(tmp_path):
    env = _env(tmp_path)
    _book(env.conn)
    csv = X.to_csv(X.dataset(env.conn, "holdings"))
    assert csv.startswith("﻿ticker,") and "KO" in csv
    out = X.export(env.conn, list(X.DATASETS), "xlsx", tmp_path / "ex", "2026-01-07")
    from openpyxl import load_workbook
    wb = load_workbook(out[0])
    assert set(wb.sheetnames) == set(X.DATASETS) and wb["nav"].max_row == 4
    files = X.export(env.conn, ["nav", "cash"], "csv", tmp_path / "ex", "d")
    assert [f.name for f in files] == ["nav_d.csv", "cash_d.csv"]
    j = json.loads(X.export(env.conn, ["nav"], "json", tmp_path / "ex", "d")[0].read_text())
    assert j["nav"][-1]["account_index"] == pytest.approx(1.0203)
    with pytest.raises(ValueError):
        X.dataset(env.conn, "secrets")


def test_export_api_and_cli(tmp_path, capsys):
    env = _env(tmp_path)
    _book(env.conn)
    c = TestClient(create_app(env.paths, PORT, spawn=lambda: None), base_url=f"http://127.0.0.1:{PORT}")
    r = c.get("/api/export/holdings.csv", headers=H)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert c.get("/api/export/all.xlsx", headers=H).content[:2] == b"PK"
    assert c.get("/api/export/../x.csv", headers=H).status_code == 404
    assert c.get("/api/export/nope.csv", headers=H).status_code == 404
    rc = cli.main(["export", "--what", "nav", "--format", "csv", "--state-dir", str(env.paths.state)])
    assert rc == 0 and "nav_" in capsys.readouterr().out


def _decision(conn, did, t, d, log):
    db.insert(conn, "decisions", {"decision_id": did, "ticker": t, "decision_date": d, "run_id": "r1",
                                  "signal": log.get("action"), "data": json.dumps(log), "created_at": db.now_iso()})


def test_validation_report(tmp_path):
    env = _env(tmp_path)
    c = env.conn
    _book(c)
    db.insert(c, "orders", {"order_id": "o1", "ticker": "KO", "side": "BUY", "qty": 100, "limit_price": 55.0,
                            "status": "filled", "created_date": "2026-01-02", "updated_at": db.now_iso()})
    db.insert(c, "fills", {"order_id": "o1", "ticker": "KO", "side": "BUY", "qty": 100, "price": 50.0, "fill_date": "2026-01-05"})
    _decision(c, "KO-1", "KO", "2026-01-02", {"action": "BUY", "gate1_not_deleted": True, "gate2_price_below_cap": True,
                                              "gate3_pass": True, "gate3_answers_file": None, "fy_last_filed": "2025-02-01",
                                              "skill_hashes": "h1"})
    _decision(c, "PEP-1", "PEP", "2026-01-03", {"action": "NO_ACTION", "stopped_at_gate": "G1", "fy_last_filed": "2026-01-03",
                                                "gate3_failures": ["Q3"], "skill_hashes": "h2"})
    _decision(c, "KO-2", "KO", "2026-01-06", {"action": "SELL_ALL", "review_type": "HOLD_REVIEW", "reason": "跌破停損線"})
    sdir = env.paths.state / "screens" / "2026-01-01"
    sdir.mkdir(parents=True)
    pd.DataFrame({"ticker": ["KO", "PEP", "MO", "X"], "lane": ["核心道", "核心道", "刪除（結構性損害／制度類）", "資料不足"]}
                 ).to_csv(sdir / "screen.csv", index=False)
    db.insert(c, "corporate_actions", {"ticker": "PEP", "action_type": "分拆", "effective_date": "2026-01-04",
                                       "status": "pending_human", "created_at": db.now_iso()})
    fa = FakeAlpaca({"KO": (50, 60), "SPY": (500, 510), "PEP": (100, 90), "MO": (40, 44)})
    r = V.build(c, env.paths, None, None, alpaca=fa)

    cm = r["compliance"]
    assert cm["n_decisions"] == 3 and cm["n_full"] == 2
    assert cm["bad_buys"] == [{"decision_id": "KO-1", "ticker": "KO", "problems": ["問卷作答檔缺"]}]
    assert cm["late_filings"] == [{"decision_id": "PEP-1", "fy_last_filed": "2026-01-03"}]
    assert cm["price_stop_sells"] == [{"decision_id": "KO-2", "ticker": "KO"}]
    assert cm["skill_hash_values"] == ["h1", "h2"]
    b = r["market"]["buys"][0]
    assert b["return"] == pytest.approx(0.2) and b["diff"] == pytest.approx(0.18)
    p = r["population"]
    assert p["groups"]["核心道"] == {"n": 1, "mean": -0.1, "median": -0.1}   # KO 已買進，不計入
    assert p["groups"]["刪除"]["mean"] == pytest.approx(0.1)
    assert r["gates"]["stopped_at"] == {"G1": 1, "未停止": 1} and r["gates"]["gate3_failures"] == {"Q3": 1}
    _decision(c, "PEP-2", "PEP", "2026-01-02", {"action": "WATCH", "gate3_pass": True})
    g = V.gates(V._logs(c, "2026-01-02", "2026-01-07"), {"PEP": -0.1}, c, "2026-01-07")
    assert g["gate3_passed"][-1] == {"ticker": "PEP", "decision_id": "PEP-2", "return": None,
                                     "noncomparable": "2026-01-04 分拆"}
    md = V.render(r)
    assert "## 5. 誠實標示" in md and "價格停損／移動停損類賣出：1 筆" in md
    mdp, jsp = V.write(r, tmp_path / "rep")
    assert mdp.is_file() and json.loads(jsp.read_text())["end"] == "2026-01-07"


def test_population_needs_period(tmp_path):
    env = _env(tmp_path)
    _book(env.conn)
    sdir = env.paths.state / "screens" / "2026-01-07"
    sdir.mkdir(parents=True)
    pd.DataFrame({"ticker": ["KO"], "lane": ["核心道"]}).to_csv(sdir / "screen.csv", index=False)
    p = V.population(env.paths, FakeAlpaca({}), "2026-01-07", "2026-01-07", set())
    assert "沒有可比較的期間" in p["note"]


def test_validation_without_prices(tmp_path):
    env = _env(tmp_path)
    with pytest.raises(ValueError):
        V.build(env.conn, env.paths, None, None)
    _book(env.conn)
    r = V.build(env.conn, env.paths, None, None, alpaca=None)
    assert "未連 Alpaca" in r["market"]["note"] and "note" in r["population"]


def test_launchd_plists(tmp_path):
    env = _env(tmp_path)
    s = L.server_plist(env.paths, 8787, tmp_path / "config.yaml")
    assert s["Label"] == L.LABEL_SERVER and s["KeepAlive"] is True
    assert s["ProgramArguments"][1:4] == ["-m", "trader.cli", "dashboard"] and "--config" in s["ProgramArguments"]
    assert "PATH" in s["EnvironmentVariables"]
    w = L.widget_plist(env.paths, 8787)
    assert w["ProgramArguments"][0].endswith("BerkshireWidget") and str(env.paths.state) in w["ProgramArguments"]
    blob = plistlib.dumps(s) + plistlib.dumps(w)
    assert b"API_KEY" not in blob and b"SECRET" not in blob


def test_validation_gate_string_pass_counts_as_true():
    """PRICE 回寫閘門 2 為字串「通過」：月報合規檢查不得誤報違規（與 sizing 判定一致）。"""
    from trader.report import validation
    from trader.portfolio import sizing
    for v in (True, "通過", "true", False, "不通過", None, "true | false"):
        assert validation._truthy(v) == sizing._truthy(v)
