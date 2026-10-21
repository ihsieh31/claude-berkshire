"""F3-CLI（A3 round3，V3-4 G2-metrics-report）：總覽待入帳入金（A3-M8C2-02）、停損否定句（A3-M8A-03）、
analyses 帶 provenance（A3-M8B-05）、觀察名單收盤（A3-M8C2-03）、PORTFOLIO 框架單另列（A3-M8C2-06）、
檢討報告列出停損賣出／晚申報／缺價代號（A3-M8C2-07）、單根 K 不算缺價（A3-M8C2-08）、registry 資料集＝非持股
（A3-M8C2-09）、排程健康：全域暫停與從未完成 daily（A3-M8C2-12）、xlsx 控制字元與公式（A3-M8E-05）。
G1／G3／G7（A3-M8A-01、M8C2-01、M8B2-01、M8B-02、M8B-03、M8B-06、M8E-01、M8C-02、M8D-06）已由 F3-LEDGER 修好，
見 test_fix_F3_LEDGER.py。全部離線。"""
from __future__ import annotations

import inspect
import io
import json
from pathlib import Path

import pandas as pd
import pytest

from test_approvals import _env
from test_dashboard import _book, _cash

from trader import db
from trader.ledger import book
from trader.paths import get_paths
from trader.pipelines import daily as DY
from trader.report import export as X
from trader.report import metrics as M
from trader.report import validation as V

APP_JS = Path(__file__).resolve().parent.parent / "trader" / "web" / "static" / "app.js"


@pytest.fixture
def conn(state_dir):
    paths = get_paths().ensure()
    c = db.connect(paths.db)
    db.migrate(c)
    yield c
    c.close()


def _decision(conn, did, t, d, log, run="r1", provenance=None):
    row = {"decision_id": did, "ticker": t, "decision_date": d, "run_id": run, "signal": log.get("action"),
           "data": json.dumps(log, ensure_ascii=False), "created_at": db.now_iso()}
    if provenance is not None:
        row["provenance"] = provenance
    db.insert(conn, "decisions", row)


class FakeAlpaca:
    """{ticker: (首, 末)}；start == end 時只回一根 K（成交日＝期末）。"""
    def __init__(self, closes):
        self.closes = closes

    def bars(self, symbols, start, end, adjustment="raw"):
        rows = []
        for t in symbols:
            if t in self.closes:
                a, b = self.closes[t]
                rows.append({"ticker": t, "date": start, "close": a})
                if end != start:
                    rows.append({"ticker": t, "date": end, "close": b})
        return pd.DataFrame(rows, columns=["ticker", "date", "close"])


# ---------------------------------------------------------------- A3-M8C2-02 待入帳入金

def test_summary_pending_deposit(tmp_path):
    env = _env(tmp_path)
    _book(env.conn)                                              # 最後淨值 2026-01-07，入金 10 萬
    _cash(env.conn, "2026-01-08", 50000.0, "deposit", "D2")      # daily 後入金、下一 daily 前
    s = M.summary(env.conn)
    assert s["deposits"] == 100000.0 and s["pending_deposits"] == 50000.0
    assert s["pnl"]["total"] == 2030.0                           # 修前 −47,970
    assert "pending_deposits" in APP_JS.read_text(encoding="utf-8")


# ---------------------------------------------------------------- A3-M8A-03 停損否定句

@pytest.mark.parametrize("text,hit", [("非因股價下跌觸發停損", False), ("並未觸發停損", False), ("不以價格停損", False),
                                      ("非價格停損", False), ("觸發停損賣出", True), ("跌破停損線", True)])
def test_stop_sell_negations(text, hit):
    assert V._stop_sell({"reason": text}) is hit


# ---------------------------------------------------------------- A3-M8B-05 provenance

def test_analyses_export_has_provenance(conn):
    _decision(conn, "KO-1", "KO", "2026-01-02", {"action": "WATCH"})
    _decision(conn, "PEP-1", "PEP", "2026-01-03", {"action": "WATCH"},
              provenance=json.dumps({"backend": "hark_queue", "effort": "max"}))
    rows = {r["decision_id"]: r for r in M.analyses(conn)}
    assert rows["KO-1"]["provenance"] is None                    # NULL＝A 機原生
    assert rows["PEP-1"]["provenance"] == {"backend": "hark_queue", "effort": "max"}
    csv = X.to_csv(X.dataset(conn, "analyses"))
    assert "provenance" in csv.splitlines()[0] and "hark_queue" in csv


# ---------------------------------------------------------------- A3-M8C2-03 觀察名單收盤

def test_watchlist_price_from_watch_closes(conn):
    book.ensure_initial_deposit(conn, 100000, "2026-01-02")
    db.insert(conn, "registry", {"ticker": "OTIS", "status": "WATCH", "data": "{}", "updated_at": db.now_iso()})
    row = book.record_nav(conn, "2026-01-02", {}, 500.0, watch_closes={"OTIS": 80.0})
    assert row["data"]["prices"] == {} and row["data"]["watch_prices"] == {"OTIS": 80.0}
    assert row["nav"] == 100000.0                                # 觀察名單不計淨值
    assert [w["price"] for w in M.watchlist(conn) if w["ticker"] == "OTIS"] == [80.0]
    assert "watch_closes=" in inspect.getsource(DY)              # daily 傳觀察名單收盤


# ---------------------------------------------------------------- A3-M8C2-06 PORTFOLIO 另列

def test_compliance_portfolio_listed_separately(tmp_path):
    env = _env(tmp_path)
    c = env.conn
    _book(c)
    _decision(c, "PEP-1", "PEP", "2026-01-03", {"action": "NO_ACTION", "stopped_at_gate": "G1"})
    _decision(c, "KO-P", "KO", "2026-01-06", {"action": "TRIM", "review_type": "PORTFOLIO", "side": "SELL",
                                             "qty": 10, "reason_one_liner": "制度類硬上限減碼"}, run="daily-2026-01-06")
    r = V.build(c, env.paths, None, None)
    cm = r["compliance"]
    assert cm["n_decisions"] == 1 and cm["n_actual"] == 1        # 修前 2／2
    assert [x["decision_id"] for x in cm["portfolio_trades"]] == ["KO-P"]
    md = V.render(r)
    assert "框架自發" in md and "KO-P" in md


# ---------------------------------------------------------------- A3-M8C2-07 render 列明細

def test_render_lists_stop_sells_late_and_missing(tmp_path):
    env = _env(tmp_path)
    c = env.conn
    _book(c)
    _decision(c, "PEP-1", "PEP", "2026-01-03", {"action": "NO_ACTION", "stopped_at_gate": "G1",
                                               "fy_last_filed": "2026-01-03"})
    _decision(c, "KO-2", "KO", "2026-01-06", {"action": "SELL_ALL", "review_type": "HOLD_REVIEW", "reason": "跌破停損線"})
    sdir = env.paths.state / "screens" / "2026-01-01"
    sdir.mkdir(parents=True)
    pd.DataFrame({"ticker": ["MO", "ZZ"], "lane": ["核心道", "核心道"]}).to_csv(sdir / "screen.csv", index=False)
    r = V.build(c, env.paths, None, None, alpaca=FakeAlpaca({"SPY": (500, 510), "MO": (40, 44)}))
    md = V.render(r)
    stop = next(x for x in md.splitlines() if x.startswith("- 價格停損"))
    late = md.split("- 使用決策日（含）之後申報的資料")[1].split("\n- ")[0]
    assert "KO-2" in md.split(stop)[1].split("\n- ")[0] and "PEP-1" in late
    assert "ZZ" in md.split("缺價格資料")[1]


# ---------------------------------------------------------------- A3-M8C2-08 單根 K 不算缺價

def test_returns_single_bar_not_missing(tmp_path):
    out, miss, short = V._returns(FakeAlpaca({"KO": (50, 50)}), ["KO", "ZZ"], "2026-01-07", "2026-01-07")
    assert out == {} and miss == ["ZZ"] and short == ["KO"]
    env = _env(tmp_path)
    c = env.conn
    _book(c)
    db.insert(c, "orders", {"order_id": "o1", "ticker": "KO", "side": "BUY", "qty": 1, "limit_price": 61.0,
                            "status": "filled", "created_date": "2026-01-06", "updated_at": db.now_iso()})
    db.insert(c, "fills", {"order_id": "o1", "ticker": "KO", "side": "BUY", "qty": 1, "price": 60.0,
                           "fill_date": "2026-01-07"})                     # 成交日＝期末
    r = V.build(c, env.paths, None, None, alpaca=FakeAlpaca({"KO": (60, 60), "SPY": (510, 510)}))
    b = r["market"]["buys"][0]
    assert b["missing"] == [] and sorted(b["short"]) == ["KO", "SPY"]
    row = next(x for x in V.render(r).splitlines() if x.startswith("| 2026-01-07 | KO"))
    assert "期間太短" in row and V.MISSING_PRICE not in row


# ---------------------------------------------------------------- A3-M8C2-09 registry 資料集＝非持股

def test_registry_dataset_labeled_non_holding(tmp_path):
    env = _env(tmp_path)
    _book(env.conn)
    assert "非持股" in X.meta(env.conn)["notes"]["registry"]
    js = APP_JS.read_text(encoding="utf-8")
    btn = js.split('href="/api/export/registry.csv"')[1][:40]
    assert "非持股" in btn


# ---------------------------------------------------------------- A3-M8C2-12 排程健康

def test_health_global_pause_and_never_ran(conn, state_dir):
    book.ensure_initial_deposit(conn, 100000, "2026-01-02")
    h = M.health(conn, state_dir)
    assert h["last_daily"] is None and "尚無完成的 daily" in (h["warning"] or "")
    book.pause(conn, None, date="2026-01-02", note="test")
    w = M.health(conn, state_dir)["warning"]
    assert "全部代號人工暫停中" in w and "1 檔暫停交易" not in w
    book.pause(conn, "KO", date="2026-01-02", note="test")
    assert "1 檔暫停交易" in M.health(conn, state_dir)["warning"]


def test_health_empty_db_no_never_ran_warning(conn, state_dir):
    assert M.health(conn, state_dir)["warning"] is None          # 全新 state（尚未入金）不吵


# ---------------------------------------------------------------- A3-M8E-05 xlsx 控制字元與公式

def test_xlsx_ctrl_char_and_formula():
    from openpyxl import load_workbook
    body = X.to_xlsx({"s": [{"a": '=HYPERLINK("http://x","y")', "b": "ab\x0ccd", "c": 1.5, "d": "-5"}]})
    ws = load_workbook(io.BytesIO(body))["s"]
    a, b, cc, d = (ws.cell(2, i) for i in range(1, 5))
    assert a.value == '=HYPERLINK("http://x","y")' and a.data_type == "s"
    assert b.value == "abcd" and cc.value == 1.5 and d.value == "-5"
